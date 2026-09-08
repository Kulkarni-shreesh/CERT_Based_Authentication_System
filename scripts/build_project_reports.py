from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
OUT = PROJECT / "docs"
IMG = OUT / "images"

REPORT_DOCX = OUT / "Cert-Based Authentication - Comprehensive Project Report.docx"
RUNBOOK_DOCX = OUT / "Cert-Based Authentication - Code and Commands Runbook.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 42, 57)
MUTED = RGBColor(92, 105, 124)
TEAL = RGBColor(15, 118, 110)


def ensure_dirs():
    OUT.mkdir(exist_ok=True)
    IMG.mkdir(exist_ok=True)


def font(size=24, bold=False):
    candidates = [
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def wrap(draw, text, fnt, width):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if draw.textbbox((0, 0), trial, font=fnt)[2] <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_box(draw, xy, title, subtitle="", fill="#FFFFFF", outline="#91A4B7"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=3)
    title_font = font(25, True)
    sub_font = font(17)
    y = y1 + 18
    for line in wrap(draw, title, title_font, x2 - x1 - 36):
        draw.text((x1 + 18, y), line, fill="#172033", font=title_font)
        y += 30
    if subtitle:
        y += 4
        for line in wrap(draw, subtitle, sub_font, x2 - x1 - 36):
            draw.text((x1 + 18, y), line, fill="#435064", font=sub_font)
            y += 23


def arrow(draw, start, end, color="#0F766E"):
    draw.line([start, end], fill=color, width=4)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) > abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        pts = [(x2, y2), (x2 - direction * 15, y2 - 9), (x2 - direction * 15, y2 + 9)]
    else:
        direction = 1 if y2 > y1 else -1
        pts = [(x2, y2), (x2 - 9, y2 - direction * 15), (x2 + 9, y2 - direction * 15)]
    draw.polygon(pts, fill=color)


def save_architecture():
    im = Image.new("RGB", (1600, 900), "#F5F7FA")
    d = ImageDraw.Draw(im)
    d.text((50, 42), "High-Level Architecture", fill="#172033", font=font(40, True))
    d.text((50, 92), "Nginx performs TLS and client certificate verification before the Python app sees the request.", fill="#52606D", font=font(21))
    boxes = {
        "browser": (70, 220, 370, 420),
        "nginx": (500, 205, 850, 435),
        "backend": (980, 220, 1290, 420),
        "db": (980, 570, 1290, 750),
        "pki": (500, 570, 850, 750),
    }
    draw_box(d, boxes["browser"], "Browser / Client", "Imports .p12 and presents a client certificate during TLS.", "#FFFFFF")
    draw_box(d, boxes["nginx"], "Nginx mTLS Gateway", "Validates certificate chain, expiry, and CRL. Adds trusted identity headers.", "#E6F6F4", "#0F766E")
    draw_box(d, boxes["backend"], "Python Backend", "Reads X-SSL-Client-* headers and maps certificate CN to an application user.", "#FFFFFF")
    draw_box(d, boxes["db"], "SQLite Database", "Stores users, employees, requests, issued certificates, sessions, notifications, and audit logs.", "#FFFFFF")
    draw_box(d, boxes["pki"], "Private PKI", "Root CA, Intermediate CA, server cert, client certs, PKCS#12 bundles, and CRLs.", "#F2F6FB", "#2E74B5")
    arrow(d, (370, 320), (500, 320))
    d.text((395, 284), "client cert", fill="#0F766E", font=font(17, True))
    arrow(d, (850, 320), (980, 320))
    d.text((888, 284), "headers", fill="#0F766E", font=font(17, True))
    arrow(d, (1135, 420), (1135, 570))
    d.text((1160, 484), "lookup / write", fill="#0F766E", font=font(17, True))
    arrow(d, (675, 570), (675, 435))
    d.text((705, 492), "trust chain + CRL", fill="#2E74B5", font=font(17, True))
    im.save(IMG / "architecture.png")


def save_mtls_flow():
    im = Image.new("RGB", (1600, 900), "#F8FAFC")
    d = ImageDraw.Draw(im)
    d.text((50, 42), "mTLS Login Flow", fill="#172033", font=font(40, True))
    steps = [
        ("1", "Client opens https://localhost/app", "Browser starts TLS connection to Nginx."),
        ("2", "Nginx asks for client certificate", "ssl_verify_client on makes the certificate mandatory."),
        ("3", "Nginx verifies certificate", "Checks CA chain, date validity, and ca-chain.crl."),
        ("4", "Decision point", "SUCCESS continues; expired, revoked, or missing cert is rejected."),
        ("5", "Backend receives trusted headers", "X-SSL-Client-Verify, subject DN, and issuer DN are forwarded."),
        ("6", "Application maps CN to user", "The backend extracts CN and queries SQLite for the user record."),
    ]
    x, y = 80, 175
    box_h = 165
    for i, title, body in steps:
        draw_box(d, (x, y, x + 410, y + box_h), f"{i}. {title}", body, "#FFFFFF", "#B8C5D6")
        if i not in ("3", "6"):
            arrow(d, (x + 410, y + 82), (x + 495, y + 82))
        x += 500
        if x > 1200:
            x = 80
            y += 230
    arrow(d, (1290, 340), (1290, 405))
    arrow(d, (490, 487), (580, 487))
    d.rounded_rectangle((82, 700, 1518, 820), radius=16, fill="#E6F6F4", outline="#0F766E", width=3)
    d.text((112, 723), "Important result", fill="#0F4F49", font=font(24, True))
    y_text = 762
    for line in wrap(d, "The Python backend does not perform cryptographic certificate validation. That job stays at the TLS gateway, and the app trusts only the local Nginx proxy boundary.", font(21), 1350):
        d.text((112, y_text), line, fill="#172033", font=font(21))
        y_text += 28
    im.save(IMG / "mtls-flow.png")


def save_cert_lifecycle():
    im = Image.new("RGB", (1600, 900), "#F8FAFC")
    d = ImageDraw.Draw(im)
    d.text((50, 42), "Certificate Lifecycle and Approval Workflow", fill="#172033", font=font(38, True))
    steps = [
        ("Employee", "submits certificate request"),
        ("Manager", "approves or rejects"),
        ("Admin", "signs approved request"),
        ("OpenSSL", "creates key, CSR, cert, and .p12"),
        ("Employee", "imports .p12 into browser"),
        ("Admin", "revokes when needed"),
        ("CRL", "is regenerated and loaded by Nginx"),
    ]
    x_positions = [70, 290, 510, 730, 950, 1170, 1390]
    y = 300
    for idx, ((title, body), x) in enumerate(zip(steps, x_positions)):
        d.ellipse((x, y, x + 120, y + 120), fill="#E6F6F4" if idx < 5 else "#FFF1F0", outline="#0F766E" if idx < 5 else "#B42318", width=4)
        label = str(idx + 1)
        bbox = d.textbbox((0, 0), label, font=font(35, True))
        d.text((x + 60 - (bbox[2] - bbox[0]) / 2, y + 38), label, fill="#172033", font=font(35, True))
        d.text((x - 30, y + 150), title, fill="#172033", font=font(22, True))
        for n, line in enumerate(wrap(d, body, font(16), 160)):
            d.text((x - 30, y + 182 + 21 * n), line, fill="#435064", font=font(16))
        if idx < len(steps) - 1:
            arrow(d, (x + 122, y + 60), (x_positions[idx + 1] - 5, y + 60), "#2E74B5")
    d.rounded_rectangle((90, 650, 1510, 795), radius=16, fill="#FFFFFF", outline="#B8C5D6", width=3)
    d.text((120, 677), "Lifecycle states in SQLite", fill="#172033", font=font(24, True))
    d.text((120, 717), "Requests: PENDING_MANAGER -> MANAGER_APPROVED -> ADMIN_SIGNED or REJECTED. Certificates: ACTIVE, EXPIRED, or REVOKED.", fill="#435064", font=font(21))
    im.save(IMG / "certificate-lifecycle.png")


def save_test_matrix():
    im = Image.new("RGB", (1600, 900), "#FFFFFF")
    d = ImageDraw.Draw(im)
    d.text((50, 42), "Security Test Results", fill="#172033", font=font(40, True))
    rows = [
        ("Valid certificate", "phase1-client", "200 OK", "SUCCESS", "#E8F5E9"),
        ("Expired certificate", "expired-client", "400 Bad Request", "FAILED: certificate has expired", "#FFF7DB"),
        ("Revoked certificate", "revoked-client", "400 Bad Request", "FAILED: certificate revoked", "#FDECEA"),
        ("No certificate", "none", "400 Bad Request", "NONE; certificate required", "#FDECEA"),
    ]
    cols = [80, 500, 850, 1160, 1520]
    headers = ["Scenario", "Certificate", "HTTP result", "Verification"]
    y = 165
    d.rounded_rectangle((70, y, 1530, y + 72), radius=10, fill="#132033")
    for i, h in enumerate(headers):
        d.text((cols[i] + 18, y + 23), h, fill="#FFFFFF", font=font(20, True))
    y += 72
    for scenario, cert, http, verify, fill in rows:
        d.rectangle((70, y, 1530, y + 115), fill=fill, outline="#D9E2EC", width=2)
        for i, text in enumerate([scenario, cert, http, verify]):
            d.text((cols[i] + 18, y + 35), text, fill="#172033", font=font(19, True if i == 0 else False))
        y += 115
    d.text((80, 720), "Evidence files are stored under evidence/ and show that Nginx blocks invalid client certificates before backend routing.", fill="#435064", font=font(22))
    im.save(IMG / "test-results.png")


def make_diagrams():
    save_architecture()
    save_mtls_flow()
    save_cert_lifecycle()
    save_test_matrix()


def set_cell(cell, text, bold=False, size=9.5):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = INK


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        shade_cell(hdr[i], "F2F4F7")
        set_cell(hdr[i], h, bold=True, size=9.2)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], str(value), size=9.1)
    if widths:
        set_table_width(table, widths)
    doc.add_paragraph()
    return table


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    for line in text.strip("\n").splitlines():
        run = p.add_run(line + "\n")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(8.6)
        run.font.color.rgb = RGBColor(40, 48, 61)
    p._p.get_or_add_pPr().append(OxmlElement("w:keepLines"))
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table, [6.3])
    cell = table.cell(0, 0)
    shade_cell(cell, "E6F6F4")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = TEAL
    r.font.size = Pt(10.5)
    p.add_run("\n" + body)
    doc.add_paragraph()


def add_figure(doc, image_name, caption):
    doc.add_picture(str(IMG / image_name), width=Inches(6.3))
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED
        run.italic = True


def setup_doc(title, subtitle, preset="standard"):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10 if preset == "standard" else 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16 if preset == "standard" else 18, 8 if preset == "standard" else 10),
        ("Heading 2", 13, BLUE, 12 if preset == "standard" else 14, 6 if preset == "standard" else 7),
        ("Heading 3", 12, DARK_BLUE, 8 if preset == "standard" else 10, 4 if preset == "standard" else 5),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = "Certificate-Based Authentication Project"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.text = "OpenSSL | Nginx mTLS | Python | SQLite"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = MUTED

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(25)
    r.font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(subtitle)
    r.font.size = Pt(13)
    r.font.color.rgb = MUTED

    meta = [
        ("Project", "Certificate-Based Authentication Lab and PKI Portal"),
        ("Prepared", "August 3, 2026"),
        ("Environment", r"C:\Users\kulka\OneDrive\Desktop\cert-based authentication"),
        ("Core stack", "OpenSSL 3.2.4, Nginx 1.30.1, Python standard library, SQLite"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        a = p.add_run(label + ": ")
        a.bold = True
        b = p.add_run(value)
        a.font.size = b.font.size = Pt(10.5)

    doc.add_paragraph()
    return doc


def build_report():
    doc = setup_doc(
        "Certificate-Based Authentication Project Report",
        "Comprehensive architecture, workflow, implementation, security behavior, testing evidence, and limitations.",
        "standard",
    )
    add_callout(doc, "Executive summary", "This project demonstrates a complete local certificate-based authentication system. A private PKI issues server and client certificates, Nginx enforces mutual TLS, a Python backend maps verified certificate identity to users, SQLite stores operational data, and CRLs prove revocation behavior.")

    doc.add_heading("1. Project Purpose and Scope", level=1)
    doc.add_paragraph("The project is a local learning and CV showcase lab for certificate-based authentication. It replaces username/password access for the protected mTLS route with browser-presented client certificates. It also expands the lab into a small PKI operations portal with request approval, certificate issuance, revocation, audit logging, and dashboards.")
    add_bullets(doc, [
        "Create a private Root CA and Intermediate CA with OpenSSL.",
        "Generate server certificates, client certificates, and browser-importable PKCS#12 bundles.",
        "Configure Nginx to require and verify client certificates.",
        "Forward verified certificate identity details to a Python backend using trusted proxy headers.",
        "Map certificate Common Name values to users in SQLite.",
        "Demonstrate expired, revoked, missing, and valid certificate outcomes with saved evidence.",
    ])

    doc.add_heading("2. Technology Stack", level=1)
    add_table(doc, ["Technology", "Role in Project"], [
        ("OpenSSL 3.2.4", "Creates CA hierarchy, keys, CSRs, certificates, PKCS#12 bundles, and CRLs."),
        ("Nginx 1.30.1", "Acts as HTTPS server, mTLS verification gateway, reverse proxy, and CRL enforcement point."),
        ("Python", "Implements the backend portal using http.server, routing, HTML rendering, session handling, subprocess calls, and SQLite access."),
        ("SQLite", "Stores users, employees, certificate requests, issued certificates, sessions, audit events, and notifications."),
        ("Browser", "Imports .p12 client certificates and presents them during the TLS handshake."),
        ("PowerShell", "Runs OpenSSL, backend, Nginx, test commands, and operational commands."),
    ], [1.5, 4.8])

    doc.add_heading("3. Architecture", level=1)
    add_figure(doc, "architecture.png", "Figure 1. The mTLS gateway owns certificate validation; the backend consumes only trusted headers from the local proxy.")
    doc.add_paragraph("Nginx is the trust boundary. The backend listens on 127.0.0.1:8081 and is intended to be reached only through Nginx. This separation keeps cryptographic certificate validation in the TLS layer while allowing the application layer to focus on user mapping, workflow, and portal behavior.")
    add_table(doc, ["Layer", "Main Files", "Responsibility"], [
        ("PKI", "pki/certs, pki/private, pki/crl, pki/*-ca-db", "Root CA, Intermediate CA, issued certificates, private keys, certificate database, and CRLs."),
        ("TLS gateway", "nginx/nginx-mtls.conf", "Requires client certificates, verifies chain/depth/CRL, and forwards verified identity headers."),
        ("Application", "app/app.py", "Handles login, mTLS result display, certificate request workflow, admin actions, downloads, and audit pages."),
        ("Data", "app/cert_auth.db", "Stores users, portal accounts, sessions, certificate lifecycle records, and evidence-oriented data."),
        ("Evidence", "evidence/*.txt and logs/*.log", "Contains test output proving accept/reject behavior for important certificate states."),
    ], [1.2, 2.0, 3.1])

    doc.add_heading("4. PKI Design", level=1)
    doc.add_paragraph("The PKI uses a Root CA as the trust anchor and an Intermediate CA for issuing operational server and client certificates. This mirrors common production PKI structure: the Root CA remains the top-level trust anchor, while the Intermediate CA signs daily-use certificates and maintains the issuing database.")
    add_table(doc, ["Certificate / Object", "Purpose", "Important Project Files"], [
        ("Root CA", "Top-level trust anchor.", "ca.crt, pki/private/ca.key, pki/root-ca-db/root-ca.cnf"),
        ("Intermediate CA", "Issues server and client certificates.", "intermediate.crt, pki/private/intermediate.key, pki/intermediate-ca-db/intermediate-ca.cnf"),
        ("CA chain", "Trust bundle used by Nginx for client verification.", "pki/certs/ca-chain.crt"),
        ("Server certificate", "Allows HTTPS on localhost/cert-auth.local.", "pki/certs/server-fullchain.crt, pki/private/server.key"),
        ("Client certificates", "Represent client identity for mTLS.", "pki/certs/*.crt, pki/private/*.key, *.p12"),
        ("CRLs", "List revoked certificates and allow Nginx to block them.", "pki/crl/intermediate.crl, pki/crl/root.crl, pki/crl/ca-chain.crl"),
    ], [1.45, 2.0, 2.85])
    doc.add_paragraph("Important extension behavior: CA certificates use keyCertSign and cRLSign; the server certificate uses serverAuth and Subject Alternative Names for localhost, 127.0.0.1, ::1, and cert-auth.local; client certificates use clientAuth.")

    doc.add_heading("5. mTLS Authentication Workflow", level=1)
    add_figure(doc, "mtls-flow.png", "Figure 2. Request handling path from browser certificate selection to backend user mapping.")
    add_numbered(doc, [
        "The browser opens https://localhost/app.",
        "Nginx requests a client certificate because ssl_verify_client is enabled.",
        "The browser presents a client certificate from the user's certificate store.",
        "Nginx validates the certificate chain against pki/certs/ca-chain.crt.",
        "Nginx checks validity dates and pki/crl/ca-chain.crl for revocation.",
        "If verification succeeds, Nginx forwards the request to 127.0.0.1:8081 with X-SSL-Client-* headers.",
        "The backend extracts the certificate Common Name and maps it to a user record.",
    ])
    add_code(doc, r'''
ssl_client_certificate "C:/Users/kulka/OneDrive/Desktop/cert-based authentication/pki/certs/ca-chain.crt";
ssl_crl "C:/Users/kulka/OneDrive/Desktop/cert-based authentication/pki/crl/ca-chain.crl";
ssl_verify_client on;
ssl_verify_depth 2;

proxy_set_header X-SSL-Client-Verify $ssl_client_verify;
proxy_set_header X-SSL-Client-DN $ssl_client_s_dn;
proxy_set_header X-SSL-Client-Issuer-DN $ssl_client_i_dn;
''')

    doc.add_heading("6. Backend and Portal Functionality", level=1)
    doc.add_paragraph("The backend is implemented in app/app.py using Python's standard-library HTTP server stack. It includes HTML rendering, form handling, session cookies, password hashing, certificate request management, OpenSSL command execution, revocation automation, and audit/notification storage.")
    add_table(doc, ["Route / Feature", "Purpose"], [
        ("/app", "Shows the mTLS login result and certificate identity mapping."),
        ("/login and /logout", "Role-account login for portal workflows."),
        ("/app/portal", "Employee certificate request and download area."),
        ("/app/manager", "Manager approval and rejection queue."),
        ("/app/admin", "Admin dashboard with signing queue, metrics, alerts, and recent activity."),
        ("/app/admin/users", "Employee/user management."),
        ("/app/admin/certificates", "Certificate inventory and revocation entry point."),
        ("/app/certificate and /app/download", "Certificate details and .p12 download."),
        ("/app/audit", "Compliance-oriented audit log view."),
    ], [2.0, 4.3])
    add_table(doc, ["Database Table", "Stores"], [
        ("users", "CN-to-user mapping for mTLS login."),
        ("employees", "Portal accounts, manager relationship, department, role, and password hash."),
        ("sessions", "Authenticated role-login sessions."),
        ("certificate_requests", "Employee requests and manager/admin decisions."),
        ("certificates", "Issued certificate records, status, file paths, serials, expiry, and revocation metadata."),
        ("audit_log", "Workflow and security-relevant action history."),
        ("notifications", "User-facing workflow notifications."),
    ], [1.8, 4.5])

    doc.add_heading("7. Certificate Lifecycle Workflow", level=1)
    add_figure(doc, "certificate-lifecycle.png", "Figure 3. Portal-driven certificate request, approval, signing, import, and revocation workflow.")
    doc.add_paragraph("The level-up portal turns the authentication lab into a lifecycle demonstration. Employees can request certificates; managers approve; administrators sign using OpenSSL; issued PKCS#12 bundles are recorded and downloadable; administrators can revoke certificates, regenerate CRLs, reload Nginx, and leave an audit trail.")

    doc.add_heading("8. Revocation Model", level=1)
    doc.add_paragraph("Revocation is implemented with Certificate Revocation Lists. Revoking a certificate does not edit the certificate itself. Instead, OpenSSL records its serial number in the Intermediate CA database and emits an updated CRL. Nginx loads the combined CRL chain and blocks certificates whose serial appears in that CRL.")
    add_code(doc, r'''
& "C:\Program Files\Git\usr\bin\openssl.exe" ca -batch `
  -config pki\intermediate-ca-db\intermediate-ca.cnf `
  -revoke pki\certs\revoked-client.crt

& "C:\Program Files\Git\usr\bin\openssl.exe" ca -batch `
  -config pki\intermediate-ca-db\intermediate-ca.cnf `
  -gencrl -crlexts crl_ext `
  -out pki\crl\intermediate.crl
''')

    doc.add_heading("9. Test Evidence", level=1)
    add_figure(doc, "test-results.png", "Figure 4. Saved evidence proves each expected mTLS security outcome.")
    add_table(doc, ["Scenario", "Certificate", "Expected and Actual Result", "Evidence File"], [
        ("Valid certificate", "phase1-client", "Accepted: HTTP/1.1 200 OK and X-SSL-Client-Verify: SUCCESS.", "evidence/01-valid-client.txt"),
        ("Expired certificate", "expired-client", "Rejected: HTTP/1.1 400 Bad Request and FAILED:certificate has expired.", "evidence/02-expired-client.txt"),
        ("Revoked certificate", "revoked-client", "Rejected: HTTP/1.1 400 Bad Request and FAILED:certificate revoked.", "evidence/03-revoked-client.txt"),
        ("No certificate", "none", "Rejected: HTTP/1.1 400 Bad Request and X-SSL-Client-Verify: NONE.", "evidence/04-no-client-cert.txt"),
    ], [1.35, 1.35, 2.8, 1.0])

    doc.add_heading("10. Security Concepts Demonstrated", level=1)
    add_bullets(doc, [
        "Public Key Infrastructure: root trust, intermediate issuing, certificate chains, serial numbers, validity periods, and extensions.",
        "Mutual TLS: both server and client identity are checked during the TLS handshake.",
        "Certificate-based identity: the application maps the verified certificate Common Name to a local user.",
        "Trust boundary design: only Nginx validates certificates; the backend trusts headers from the local proxy boundary.",
        "Revocation: CRLs let the system disable a certificate before its expiration date.",
        "Auditability: operational actions such as requests, approvals, signing, and revocation are logged.",
    ])

    doc.add_heading("11. Limitations and Production Hardening", level=1)
    add_table(doc, ["Current Lab Limitation", "Production Hardening Idea"], [
        ("Local files hold CA private keys.", "Use secure key storage, offline Root CA procedures, restricted permissions, or an HSM."),
        ("CN is used as a simple identity key.", "Use SAN email/URI or a stable certificate subject mapping policy."),
        ("CRL is local and manually regenerated/reloaded in the lab.", "Publish CRLs or use OCSP/OCSP stapling where appropriate."),
        ("Backend trusts proxy headers by deployment design.", "Block direct backend exposure and strip inbound spoofed identity headers at the edge."),
        ("SQLite is local and single-node.", "Use a managed database and migration framework for multi-user production use."),
        ("Demo passwords and local bundle passwords are documented.", "Use secret management, enrollment controls, and per-user delivery channels."),
    ], [2.3, 4.0])

    doc.add_heading("12. Interview and CV Summary", level=1)
    add_callout(doc, "CV-ready summary", "Built a certificate-based authentication lab using OpenSSL, Nginx mTLS, Python, and SQLite. Implemented a private CA chain, browser-importable client certificates, Nginx client certificate enforcement, backend CN-to-user mapping, CRL-based revocation, a PKI request/approval/signing portal, audit logging, and evidence for valid, expired, revoked, and missing certificate scenarios.")

    doc.add_heading("13. Project Folder Map", level=1)
    add_code(doc, r'''
cert-based authentication/
  app/                 Python backend and SQLite database
  evidence/            Saved mTLS test outputs
  logs/                Nginx mTLS access and error logs
  nginx/               mTLS gateway configuration
  pki/                 CA databases, certs, keys, CSRs, CRLs, extension configs
  tools/nginx-1.30.1/  Local Nginx runtime
  www/                 Static landing page
  *.p12                Browser-importable client certificate bundles
''')
    doc.save(REPORT_DOCX)


def build_runbook():
    doc = setup_doc(
        "Code and Commands Runbook",
        "Exact commands, files, workflows, import methods, and operational procedures for running the project.",
        "compact",
    )
    add_callout(doc, "How to use this guide", "Use this document when you need to start the lab, import certificates, test mTLS, add users, issue new certificates, revoke certificates, or explain the implementation from code.")

    doc.add_heading("1. Prerequisites and Locations", level=1)
    add_table(doc, ["Item", "Value"], [
        ("Project folder", r"C:\Users\kulka\OneDrive\Desktop\cert-based authentication"),
        ("OpenSSL executable", r"C:\Program Files\Git\usr\bin\openssl.exe"),
        ("Nginx executable", r".\tools\nginx-1.30.1\nginx.exe"),
        ("Backend", r"app\app.py on 127.0.0.1:8081"),
        ("HTTPS entry point", r"https://localhost/app"),
        ("Database", r"app\cert_auth.db"),
    ], [1.8, 4.5])
    add_code(doc, r'''
cd "C:\Users\kulka\OneDrive\Desktop\cert-based authentication"
& "C:\Program Files\Git\usr\bin\openssl.exe" version
''')

    doc.add_heading("2. Start, Reload, and Stop the Project", level=1)
    doc.add_heading("Start the Python backend", level=2)
    add_code(doc, r'''
python app\app.py
''')
    doc.add_heading("Start Nginx", level=2)
    add_code(doc, r'''
.\tools\nginx-1.30.1\nginx.exe `
  -p "C:\Users\kulka\OneDrive\Desktop\cert-based authentication" `
  -c "nginx\nginx-mtls.conf"
''')
    doc.add_heading("Reload Nginx after config or CRL changes", level=2)
    add_code(doc, r'''
.\tools\nginx-1.30.1\nginx.exe `
  -p "C:\Users\kulka\OneDrive\Desktop\cert-based authentication" `
  -c "nginx\nginx-mtls.conf" `
  -s reload
''')
    doc.add_heading("Stop Nginx", level=2)
    add_code(doc, r'''
.\tools\nginx-1.30.1\nginx.exe `
  -p "C:\Users\kulka\OneDrive\Desktop\cert-based authentication" `
  -c "nginx\nginx-mtls.conf" `
  -s stop
''')

    doc.add_heading("3. Browser Certificate Import", level=1)
    add_table(doc, ["Bundle", "Purpose", "Password"], [
        ("client.p12", "Valid phase1-client demo certificate", "clientpass"),
        ("alice-client.p12", "Valid Alice demo certificate", "Use the password set when generated"),
        ("expired-client.p12", "Expired demo certificate", "expiredpass"),
        ("revoked-client.p12", "Revoked demo certificate", "revokedpass"),
        ("Portal-issued .p12", "Generated from admin signing workflow", "portalpass"),
    ], [1.6, 3.0, 1.4])
    add_numbered(doc, [
        "Double-click the .p12 file or open the browser/Windows certificate import wizard.",
        "Choose Current User when asked where to import.",
        "Enter the .p12 password.",
        "Place the certificate in the Personal certificate store when prompted.",
        "Open https://localhost/app and select the client certificate in the browser picker.",
    ])
    add_callout(doc, "Browser behavior note", "Browsers may hide expired certificates from the certificate picker. Use the OpenSSL s_client tests in this guide when you need clear evidence for expired-client behavior.")

    doc.add_heading("4. Portal Accounts", level=1)
    add_table(doc, ["Role", "Email", "Password", "Main Page"], [
        ("Employee", "john@company.com", "password123", "/app/portal"),
        ("Manager", "bob@company.com", "password123", "/app/manager"),
        ("Admin", "charlie@company.com", "password123", "/app/admin"),
    ], [1.1, 2.0, 1.4, 1.8])

    doc.add_heading("5. Add a New Certificate Through the Portal", level=1)
    add_numbered(doc, [
        "Start the backend and Nginx.",
        "Open https://localhost/login.",
        "Sign in as an employee, for example john@company.com with password123.",
        "Open the Employee Portal and submit a certificate request.",
        "Sign in as manager bob@company.com and approve the request from /app/manager.",
        "Sign in as admin charlie@company.com and click Sign & Issue from /app/admin.",
        "Download the generated .p12 from the certificate details or employee portal.",
        "Import the .p12 into the browser using the portal password: portalpass.",
        "Open https://localhost/app and test the new certificate if the CN is mapped for mTLS login behavior.",
    ])

    doc.add_heading("6. Add a New Client Certificate Manually", level=1)
    doc.add_paragraph("Use this when you want to create a client certificate outside the portal. Replace new-client with the desired Common Name.")
    add_code(doc, r'''
$cn = "new-client"

& "C:\Program Files\Git\usr\bin\openssl.exe" genrsa `
  -out "pki\private\$cn.key" 2048

& "C:\Program Files\Git\usr\bin\openssl.exe" req -new `
  -key "pki\private\$cn.key" `
  -out "pki\csr\$cn.csr" `
  -subj "/C=IN/ST=Karnataka/L=Bengaluru/O=Cert Based Auth Lab/OU=Manual/CN=$cn"

& "C:\Program Files\Git\usr\bin\openssl.exe" ca -batch `
  -config pki\intermediate-ca-db\intermediate-ca.cnf `
  -extensions client_cert `
  -days 825 `
  -notext `
  -md sha256 `
  -in "pki\csr\$cn.csr" `
  -out "pki\certs\$cn.crt"

& "C:\Program Files\Git\usr\bin\openssl.exe" pkcs12 -export `
  -inkey "pki\private\$cn.key" `
  -in "pki\certs\$cn.crt" `
  -certfile "pki\certs\ca-chain.crt" `
  -out "$cn.p12" `
  -passout pass:portalpass
''')
    doc.add_paragraph("If the new certificate should authenticate to the /app mTLS user mapping route, add a matching row in the users table with the same CN.")
    add_code(doc, r'''
python -c "import sqlite3; c=sqlite3.connect('app/cert_auth.db'); c.execute(\"insert or replace into users (cn, username, display_name, role) values (?, ?, ?, ?)\", ('new-client','new_client','New Client','student')); c.commit(); c.close()"
''')

    doc.add_heading("7. Revoke a Certificate", level=1)
    doc.add_heading("Preferred UI method", level=2)
    add_numbered(doc, [
        "Sign in as admin at https://localhost/login.",
        "Open /app/admin/certificates.",
        "Click Revoke for an ACTIVE certificate.",
        "Choose a reason and submit notes.",
        "The backend revokes the certificate, regenerates the Intermediate CRL, rebuilds ca-chain.crl, reloads Nginx, and writes audit/notification records.",
    ])
    doc.add_heading("Manual command method", level=2)
    add_code(doc, r'''
& "C:\Program Files\Git\usr\bin\openssl.exe" ca -batch `
  -config pki\intermediate-ca-db\intermediate-ca.cnf `
  -revoke pki\certs\new-client.crt

& "C:\Program Files\Git\usr\bin\openssl.exe" ca -batch `
  -config pki\intermediate-ca-db\intermediate-ca.cnf `
  -gencrl `
  -crlexts crl_ext `
  -out pki\crl\intermediate.crl

Get-Content pki\crl\intermediate.crl, pki\crl\root.crl |
  Set-Content pki\crl\ca-chain.crl

.\tools\nginx-1.30.1\nginx.exe `
  -p "C:\Users\kulka\OneDrive\Desktop\cert-based authentication" `
  -c "nginx\nginx-mtls.conf" `
  -s reload
''')

    doc.add_heading("8. Test Commands", level=1)
    add_code(doc, r'''
$req = "GET /app HTTP/1.1`r`nHost: localhost`r`nAccept: application/json`r`nConnection: close`r`n`r`n"

# Valid certificate
$req | & "C:\Program Files\Git\usr\bin\openssl.exe" s_client `
  -connect localhost:443 `
  -CAfile "pki\certs\ca.crt" `
  -cert "pki\certs\client.crt" `
  -key "pki\private\client.key" `
  -quiet

# Expired certificate
$req | & "C:\Program Files\Git\usr\bin\openssl.exe" s_client `
  -connect localhost:443 `
  -CAfile "pki\certs\ca.crt" `
  -cert "pki\certs\expired-client.crt" `
  -key "pki\private\expired-client.key" `
  -quiet

# Revoked certificate
$req | & "C:\Program Files\Git\usr\bin\openssl.exe" s_client `
  -connect localhost:443 `
  -CAfile "pki\certs\ca.crt" `
  -cert "pki\certs\revoked-client.crt" `
  -key "pki\private\revoked-client.key" `
  -quiet

# No certificate
$req | & "C:\Program Files\Git\usr\bin\openssl.exe" s_client `
  -connect localhost:443 `
  -CAfile "pki\certs\ca.crt" `
  -quiet
''')
    add_table(doc, ["Expected Case", "Expected Output"], [
        ("Valid certificate", "HTTP/1.1 200 OK and X-SSL-Client-Verify: SUCCESS."),
        ("Expired certificate", "HTTP/1.1 400 Bad Request and FAILED:certificate has expired."),
        ("Revoked certificate", "HTTP/1.1 400 Bad Request and FAILED:certificate revoked."),
        ("No certificate", "HTTP/1.1 400 Bad Request and X-SSL-Client-Verify: NONE."),
    ], [1.8, 4.5])

    doc.add_heading("9. Code Map", level=1)
    add_table(doc, ["Function / Area", "What it does"], [
        ("parse_cn(subject_dn)", "Extracts the Common Name from the Nginx-provided subject DN."),
        ("init_db()", "Creates and seeds users, employees, sessions, requests, certificates, audit log, and notifications."),
        ("create_session / get_session_user", "Creates and validates role-login sessions."),
        ("create_request()", "Creates employee certificate requests and manager notifications."),
        ("approve_request()", "Records manager approval or rejection decisions."),
        ("issue_certificate()", "Runs OpenSSL to create key, CSR, certificate, .p12 bundle, and database certificate record."),
        ("regenerate_crl_and_reload()", "Rebuilds CRL files and reloads Nginx after revocation."),
        ("revoke_certificate()", "Marks certificate as revoked, calls OpenSSL revoke, updates CRL, reloads Nginx, and audits."),
        ("render_login_result()", "Displays mTLS identity and user mapping result for /app."),
        ("RequestHandler.do_GET / do_POST", "Routes pages and form actions."),
    ], [2.2, 4.1])

    doc.add_heading("10. Important Configuration Lines", level=1)
    add_code(doc, r'''
# nginx/nginx-mtls.conf
listen 443 ssl;
ssl_certificate      ".../pki/certs/server-fullchain.crt";
ssl_certificate_key  ".../pki/private/server.key";
ssl_client_certificate ".../pki/certs/ca-chain.crt";
ssl_crl ".../pki/crl/ca-chain.crl";
ssl_verify_client on;
ssl_verify_depth 2;
proxy_pass http://127.0.0.1:8081;
''')

    doc.add_heading("11. Troubleshooting", level=1)
    add_table(doc, ["Symptom", "Likely Cause", "Fix"], [
        ("openssl is not recognized", "OpenSSL is not on PATH.", 'Use & "C:\\Program Files\\Git\\usr\\bin\\openssl.exe" explicitly.'),
        ("Browser does not show expired-client", "Browser filters expired personal certificates.", "Use OpenSSL s_client test for expired evidence."),
        ("400 No required SSL certificate was sent", "No client certificate was selected or available.", "Import .p12 into Personal store and reopen the browser."),
        ("FAILED:certificate revoked", "Certificate serial is present in CRL.", "Use a different active certificate or remove/reissue as part of a controlled lab reset."),
        ("Backend works directly on 8081 but HTTPS fails", "Nginx not started or config/reload issue.", "Start or reload Nginx with nginx-mtls.conf."),
        ("New cert signs but /app does not map user", "CN not present in users table.", "Insert a matching CN row in SQLite users table."),
    ], [1.8, 2.0, 2.5])

    doc.add_heading("12. Demo Checklist", level=1)
    add_bullets(doc, [
        "Show pki/certs/ca-chain.crt and explain Root CA plus Intermediate CA.",
        "Show nginx/nginx-mtls.conf and point to ssl_verify_client, ssl_client_certificate, ssl_crl, and proxy_set_header lines.",
        "Show app/app.py parse_cn and render_login_result behavior.",
        "Open https://localhost/app with a valid browser certificate.",
        "Show saved evidence for valid, expired, revoked, and no-certificate cases.",
        "Use the portal to submit, approve, sign, download, and revoke a certificate.",
        "Show /app/audit to prove workflow actions are recorded.",
    ])

    doc.save(RUNBOOK_DOCX)


def main():
    ensure_dirs()
    make_diagrams()
    build_report()
    build_runbook()
    print(REPORT_DOCX)
    print(RUNBOOK_DOCX)


if __name__ == "__main__":
    main()
